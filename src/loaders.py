"""Leitura dos documentos: cada formato vira uma lista de trechos {text, location}."""
import csv
import json
from pathlib import Path

MAX_SEG = 800  # tamanho aproximado (caracteres) dos grupos de linhas de tabelas


# ---------- utilitários ----------
def _fmt_row(header, values):
    """Linha de tabela -> 'coluna: valor | coluna: valor' (o cabeçalho vai junto em cada linha)."""
    pairs = []
    for h, v in zip(header, values):
        v = "" if v is None else str(v).strip()
        if v:
            pairs.append(f"{h}: {v}")
    return " | ".join(pairs)


def _group_rows(rows, prefix):
    """Agrupa linhas de tabela em trechos de ~MAX_SEG caracteres."""
    segs, buf, size, first = [], [], 0, 1
    for i, row in enumerate(rows, start=1):
        if buf and size + len(row) > MAX_SEG:
            segs.append({"text": "\n".join(buf), "location": f"{prefix}, linhas {first}-{i - 1}"})
            buf, size, first = [], 0, i
        buf.append(row)
        size += len(row) + 1
    if buf:
        segs.append({"text": "\n".join(buf), "location": f"{prefix}, linhas {first}-{len(rows)}"})
    return segs


def _table_segments(table_rows, prefix):
    """table_rows: lista de listas, primeira linha = cabeçalho."""
    table_rows = [r for r in table_rows if any(str(c or "").strip() for c in r)]
    if len(table_rows) < 2:
        return []
    header = [str(h).strip() if h not in (None, "") else f"coluna{i + 1}" for i, h in enumerate(table_rows[0])]
    rows = [_fmt_row(header, r) for r in table_rows[1:]]
    return _group_rows([r for r in rows if r], prefix)


# ---------- loaders por formato ----------
def load_pdf(path):
    from pypdf import PdfReader

    segs = []
    for n, page in enumerate(PdfReader(path).pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            segs.append({"text": text, "location": f"página {n}"})
    return segs


def load_docx(path):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    lines = []
    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            p = Paragraph(child, doc)
            if p.text.strip():
                is_heading = p.style is not None and p.style.name.lower().startswith(("heading", "título", "titulo"))
                lines.append(("# " if is_heading else "") + p.text.strip())
        elif child.tag.endswith("}tbl"):
            table = [[c.text for c in row.cells] for row in Table(child, doc).rows]
            for seg in _table_segments(table, "tabela"):
                lines.append(seg["text"])
    return [{"text": "\n".join(lines), "location": "documento"}]


def load_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    segs = []
    for ws in wb.worksheets:
        segs += _table_segments([list(r) for r in ws.iter_rows(values_only=True)], f"aba '{ws.title}'")
    return segs


def load_csv(path):
    text = Path(path).read_text(encoding="utf-8-sig")
    try:
        delimiter = csv.Sniffer().sniff(text[:2000], delimiters=",;\t").delimiter
    except csv.Error:
        delimiter = ","
    return _table_segments(list(csv.reader(text.splitlines(), delimiter=delimiter)), "tabela")


def load_pptx(path):
    from pptx import Presentation

    segs = []
    for n, slide in enumerate(Presentation(path).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                parts += [" | ".join(c.text for c in row.cells) for row in shape.table.rows]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Notas: " + slide.notes_slide.notes_text_frame.text.strip())
        if parts:
            segs.append({"text": "\n".join(parts), "location": f"slide {n}"})
    return segs


def load_md(path):
    """Markdown/TXT: divide por títulos (#) para manter o contexto de cada seção."""
    segs, title, buf = [], "início", []

    def flush():
        if "".join(buf).strip():
            segs.append({"text": "\n".join(buf), "location": f"seção '{title}'"})

    for line in Path(path).read_text(encoding="utf-8").splitlines():
        if line.startswith("#"):
            flush()
            title, buf = line.lstrip("# ").strip(), [line]
        else:
            buf.append(line)
    flush()
    return segs


def load_json(path):
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if isinstance(data, list) and data and all(isinstance(d, dict) for d in data):
        return _group_rows([json.dumps(d, ensure_ascii=False) for d in data], "registros")
    return [{"text": json.dumps(data, ensure_ascii=False, indent=2), "location": "json"}]


def load_html(path):
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(Path(path).read_text(encoding="utf-8"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    for tr in soup.find_all("tr"):  # mantém cada linha de tabela numa linha só
        cells = [c.get_text(" ", strip=True) for c in tr.find_all(["th", "td"])]
        tr.replace_with(soup.new_string(" | ".join(cells) + "\n"))
    title = soup.title.get_text(strip=True) if soup.title else "página"
    return [{"text": soup.get_text("\n", strip=True), "location": title}]


LOADERS = {
    ".pdf": load_pdf, ".docx": load_docx, ".xlsx": load_xlsx, ".csv": load_csv,
    ".pptx": load_pptx, ".md": load_md, ".txt": load_md, ".json": load_json,
    ".html": load_html, ".htm": load_html,
}
SUPPORTED = set(LOADERS)


def load_file(path):
    return LOADERS[Path(path).suffix.lower()](path)
