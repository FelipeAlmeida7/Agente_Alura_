"""Gera documentos fictícios da "Nova Aurora Tecnologia" em data/docs (PDF, Word, Excel, PowerPoint, MD, CSV, JSON, HTML).

Uso:  pip install -r requirements-dev.txt && python scripts/generate_samples.py
"""
import json
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).resolve().parent.parent / "data" / "docs"


def path(*parts):
    p = OUT.joinpath(*parts)
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


# ------------------------------------------------------------------ PDF
def make_pdf(file, title, blocks):
    """blocks: lista de ('h', texto) | ('p', texto) | ('b', [itens])."""
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 12)]
    for kind, content in blocks:
        if kind == "h":
            story.append(Paragraph(content, styles["Heading2"]))
        elif kind == "p":
            story.append(Paragraph(content, styles["Normal"]))
        else:
            story += [Paragraph(f"- {i}", styles["Normal"]) for i in content]
        story.append(Spacer(1, 8))
    SimpleDocTemplate(str(file), pagesize=A4, title=title).build(story)


make_pdf(path("rh", "politica_ferias.pdf"), "Política de Férias - Nova Aurora Tecnologia", [
    ("h", "1. Direito a férias"),
    ("p", "Após 12 meses de trabalho (período aquisitivo), o colaborador CLT tem direito a 30 dias corridos de férias."),
    ("h", "2. Fracionamento"),
    ("p", "As férias podem ser divididas em até 3 períodos. Um deles deve ter no mínimo 14 dias corridos e os demais "
          "no mínimo 5 dias corridos cada."),
    ("h", "3. Solicitação e aprovação"),
    ("p", "A solicitação é feita no portal RH+ com no mínimo 30 dias de antecedência. O gestor tem 5 dias úteis "
          "para aprovar ou sugerir outra data."),
    ("h", "4. Abono pecuniário"),
    ("p", "O colaborador pode converter até 10 dias de férias em abono pecuniário (venda de férias), solicitando "
          "até 15 dias antes do fim do período aquisitivo."),
    ("h", "5. Pagamento"),
    ("p", "O pagamento das férias, com o adicional de 1/3, é feito até 2 dias antes do início do descanso."),
])

make_pdf(path("legal", "politica_lgpd_e_nda.pdf"), "LGPD e Confidencialidade (NDA) - Nova Aurora Tecnologia", [
    ("h", "1. Encarregado de dados (DPO)"),
    ("p", "A encarregada de proteção de dados é Marina Duarte, contato: dpo@novaaurora.example."),
    ("h", "2. Incidentes de segurança"),
    ("p", "Qualquer suspeita de vazamento de dados pessoais deve ser comunicada ao DPO em até 24 horas. "
          "A comunicação à ANPD e aos titulares afetados é feita em até 3 dias úteis, conforme avaliação do DPO."),
    ("h", "3. Retenção de dados"),
    ("b", ["Dados de clientes: 5 anos após o término do contrato.",
           "Currículos de candidatos não contratados: 6 meses.",
           "Solicitações de titulares (acesso, correção, exclusão) são respondidas em até 15 dias."]),
    ("h", "4. Acordo de confidencialidade (NDA)"),
    ("p", "Todo colaborador assina o NDA na admissão. O sigilo vale durante o vínculo e por 3 anos após o seu "
          "encerramento. São confidenciais: código-fonte, base de clientes e preços negociados. É proibido "
          "compartilhar informações confidenciais por e-mail pessoal ou por ferramentas de IA não homologadas."),
])

# ------------------------------------------------------------------ DOCX
doc = Document()
doc.add_heading("Benefícios e Onboarding", level=1)
doc.add_paragraph("Resumo dos benefícios oferecidos a todos os colaboradores CLT da Nova Aurora Tecnologia.")
doc.add_heading("Benefícios", level=2)
rows = [("Benefício", "Valor", "Regra"),
        ("Vale-refeição/alimentação", "R$ 45,00 por dia útil", "Cartão flexível, crédito no 1º dia útil do mês"),
        ("Plano de saúde", "Empresa paga 80% do titular", "Dependentes têm desconto de 30% na mensalidade"),
        ("Plano odontológico", "100% pago pela empresa", "Apenas para o titular"),
        ("Auxílio home office", "R$ 150,00 por mês", "Para quem trabalha em regime híbrido ou remoto"),
        ("Auxílio academia", "50% da mensalidade, até R$ 100,00", "Mediante comprovante"),
        ("Auxílio-creche", "R$ 600,00 por mês por filho", "Para filhos de até 5 anos")]
t = doc.add_table(rows=len(rows), cols=3)
t.style = "Table Grid"
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        t.cell(i, j).text = v
doc.add_heading("Onboarding de novos colaboradores", level=2)
for item in ["A primeira semana é de integração, com apresentação da empresa e das equipes.",
             "Cada novo colaborador recebe um buddy (mentor) que o acompanha nos primeiros 90 dias.",
             "Os acessos aos sistemas são liberados em até 2 dias úteis após a admissão.",
             "As avaliações de experiência acontecem aos 45 e aos 90 dias."]:
    doc.add_paragraph(item, style="List Bullet")
doc.save(path("rh", "beneficios_e_onboarding.docx"))

# ------------------------------------------------------------------ Markdown
path("financeiro", "politica_despesas.md").write_text("""# Política de Despesas e Reembolsos

## Limites de viagem
- Alimentação: até R$ 120,00 por dia.
- Hospedagem: até R$ 450,00 por diária.
- Passagens aéreas em classe econômica, compradas com no mínimo 14 dias de antecedência.
- Táxi e aplicativos de transporte são permitidos quando não houver alternativa mais econômica.

## Alçadas de aprovação
- Até R$ 1.000,00: gestor direto.
- De R$ 1.000,01 a R$ 5.000,00: gestor direto e diretor da área.
- Acima de R$ 5.000,00: diretoria financeira (CFO).

## Reembolso
- A despesa deve ser lançada no sistema Expense em até 15 dias corridos, com nota fiscal.
- O pagamento é feito em até 10 dias úteis após a aprovação.

## Despesas não reembolsáveis
Bebidas alcoólicas, multas de trânsito e despesas pessoais.
""", encoding="utf-8")

path("pesquisa_desenvolvimento", "business_case_assistente_ia.md").write_text("""# Business Case: Assistente de IA para atendimento

## Contexto
O volume de chamados de suporte cresce cerca de 3% ao mês. A proposta é um assistente de IA que responda \
dúvidas frequentes dos clientes antes da abertura de chamado.

## Análise de concorrentes (pesquisa interna)
- Alfa Suporte: R$ 8.500 por mês, sem integração com a nossa base de conhecimento.
- BotNorte: R$ 6.200 por mês, atendimento apenas em horário comercial.
- Helpix: R$ 11.000 por mês, com integração completa, porém dados hospedados fora do Brasil.

## Números do projeto
- Investimento inicial: R$ 300 mil.
- Custo de operação: R$ 20 mil por mês.
- Economia estimada: R$ 45 mil por mês com a redução de 25% dos chamados.
- Payback estimado: 12 meses.

## Recomendação
Aprovar um piloto de 90 dias com 20% da base de clientes e avaliar CSAT e taxa de resolução.
""", encoding="utf-8")

# ------------------------------------------------------------------ Excel
wb = Workbook()
ws = wb.active
ws.title = "DRE 2025"
ws.append(["Item (valores em R$ mil)", "1T25", "2T25", "3T25", "4T25", "Total 2025"])
dre = [("Receita Bruta", [4200, 4550, 4900, 5350]),
       ("(-) Deduções (impostos)", [-546, -592, -637, -696]),
       ("Receita Líquida", [3654, 3958, 4263, 4654]),
       ("(-) Custos dos serviços prestados", [-1450, -1520, -1600, -1710]),
       ("Lucro Bruto", [2204, 2438, 2663, 2944]),
       ("(-) Despesas operacionais", [-1350, -1400, -1480, -1560]),
       ("EBITDA", [854, 1038, 1183, 1384]),
       ("(-) Depreciação e resultado financeiro", [-210, -215, -220, -230]),
       ("Lucro antes do IR/CSLL", [644, 823, 963, 1154]),
       ("(-) IR e CSLL", [-219, -280, -327, -392]),
       ("Lucro Líquido", [425, 543, 636, 762])]
for name, q in dre:
    ws.append([name, *q, sum(q)])
wb.save(path("financeiro", "dre_2025.xlsx"))

wb = Workbook()
ws = wb.active
ws.title = "Resumo"
ws.append(["Tópico", "Detalhe"])
ws.append(["Norma", "ISO 9001:2015"])
ws.append(["Auditoria", "Auditoria externa realizada em outubro de 2025"])
ws.append(["Resultado", "5 não conformidades: 2 maiores e 3 menores"])
ws.append(["Recertificação", "Prevista para março de 2026"])
ws2 = wb.create_sheet("Não conformidades")
ws2.append(["ID", "Área", "Não conformidade", "Gravidade", "Responsável", "Prazo", "Status"])
for r in [("NC-01", "Suporte", "Registros de treinamento incompletos", "Menor", "Carla Mendes", "2025-11-30", "Concluída"),
          ("NC-02", "Desenvolvimento", "Revisão de código não documentada em 3 projetos", "Maior", "Rafael Lima", "2025-12-15", "Em andamento"),
          ("NC-03", "Compras", "Avaliação anual de fornecedores atrasada", "Menor", "Juliana Prado", "2026-01-31", "Em andamento"),
          ("NC-04", "Infraestrutura", "Backup sem teste de restauração semestral", "Maior", "Diego Rocha", "2025-12-01", "Atrasada"),
          ("NC-05", "RH", "Descrição de cargos desatualizada", "Menor", "Ana Souza", "2026-02-28", "Aberta")]:
    ws2.append(list(r))
wb.save(path("qualidade", "auditoria_iso9001_2025.xlsx"))

# ------------------------------------------------------------------ PowerPoint
prs = Presentation()


def slide(title, lines, notes=None):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    tf = s.placeholders[1].text_frame
    tf.text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().text = line
    if notes:
        s.notes_slide.notes_text_frame.text = notes


s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = "OKRs 2026 - Nova Aurora Tecnologia"
s.placeholders[1].text = "Planejamento estratégico anual"
slide("Objetivo 1: Crescer a receita recorrente", [
    "KR1: MRR de R$ 1,6 mi para R$ 2,08 mi (+30%)",
    "KR2: Lançar o plano Enterprise até junho",
    "KR3: Fechar 12 contratos do plano Business",
    "Responsável: Diretoria Comercial"], "Meta anual; revisão trimestral.")
slide("Objetivo 2: Encantar os clientes", [
    "KR1: NPS de 45 para 60",
    "KR2: Tempo de primeira resposta abaixo de 1 hora",
    "KR3: CSAT de 92%",
    "Responsável: Diretoria de Sucesso do Cliente"])
slide("Objetivo 3: Operar com eficiência", [
    "KR1: Churn mensal de 6% para 4%",
    "KR2: Reduzir em 15% o custo de infraestrutura",
    "KR3: 100% dos colaboradores treinados em LGPD até março",
    "Responsável: Diretoria de Operações"])
prs.save(path("estrategico", "okrs_2026.pptx"))

# ------------------------------------------------------------------ CSV
path("marketing_comercial", "tabela_precos.csv").write_text(
    "plano;preco_mensal_brl;usuarios_incluidos;suporte;desconto_plano_anual\n"
    "Start;99,00;3;E-mail (resposta em 2 dias úteis);10%\n"
    "Pro;299,00;15;E-mail e chat (resposta em 8 horas);15%\n"
    "Business;899,00;50;Chat prioritário e gerente de conta (resposta em 2 horas);20%\n"
    "Enterprise;Sob consulta;Ilimitado;Suporte 24x7 dedicado;Negociado em contrato\n", encoding="utf-8")

# ------------------------------------------------------------------ JSON
api = {
    "nome": "API de Clientes",
    "versao": "v2",
    "base_url": "https://api.novaaurora.example/v2",
    "autenticacao": {"tipo": "Bearer token (OAuth2 client credentials)", "expiracao_token_minutos": 60},
    "limite_requisicoes": "100 requisições por minuto por chave de API",
    "endpoints": [
        {"metodo": "GET", "caminho": "/clientes", "descricao": "Lista clientes com paginação (máximo de 100 por página)"},
        {"metodo": "GET", "caminho": "/clientes/{id}", "descricao": "Detalha um cliente"},
        {"metodo": "POST", "caminho": "/clientes", "descricao": "Cria um cliente. Campos obrigatórios: razao_social, cnpj, email"},
        {"metodo": "DELETE", "caminho": "/clientes/{id}", "descricao": "Anonimiza os dados do cliente conforme a LGPD"},
    ],
    "codigos_erro": {"401": "Token inválido ou expirado", "404": "Cliente não encontrado", "429": "Limite de requisições excedido"},
    "responsavel": "Time de Plataforma (canal #plataforma)",
}
path("dados_sistemas", "api_clientes.json").write_text(json.dumps(api, ensure_ascii=False, indent=2), encoding="utf-8")

# ------------------------------------------------------------------ HTML
path("operacional", "sla_suporte.html").write_text("""<!DOCTYPE html>
<html lang="pt-BR"><head><meta charset="utf-8"><title>Procedimento de Atendimento e SLA de Suporte</title></head>
<body>
<h1>Procedimento de Atendimento e SLA de Suporte</h1>
<p>O suporte atende 24x7 chamados críticos. Os demais são atendidos em horário comercial (8h às 18h, horário de Brasília).</p>
<h2>Tabela de SLA</h2>
<table>
<tr><th>Prioridade</th><th>Descrição</th><th>Primeira resposta</th><th>Resolução</th></tr>
<tr><td>Crítica</td><td>Sistema fora do ar para todos os clientes</td><td>30 minutos</td><td>4 horas</td></tr>
<tr><td>Alta</td><td>Funcionalidade principal indisponível</td><td>2 horas</td><td>1 dia útil</td></tr>
<tr><td>Média</td><td>Falha parcial com solução de contorno</td><td>8 horas</td><td>3 dias úteis</td></tr>
<tr><td>Baixa</td><td>Dúvidas e pedidos de melhoria</td><td>1 dia útil</td><td>5 dias úteis</td></tr>
</table>
<h2>Fluxo de escalonamento</h2>
<ol>
<li>Registrar o chamado na plataforma Helpdesk e classificar a prioridade.</li>
<li>O nível 1 (N1) tenta resolver em até 30 minutos.</li>
<li>Sem solução, escalar para o nível 2 (N2), do time de engenharia.</li>
<li>Em chamados críticos, acionar o plantonista pelo canal #incidentes e avisar o gerente de plantão.</li>
<li>Após incidentes críticos, publicar o relatório pós-incidente em até 3 dias úteis.</li>
</ol>
</body></html>
""", encoding="utf-8")

path("comunicacao_interna", "newsletter_setembro_2026.html").write_text("""<!DOCTYPE html>
<html lang="pt-BR"><head><meta charset="utf-8"><title>Nova Aurora News - Setembro de 2026</title></head>
<body>
<h1>Nova Aurora News - Setembro de 2026</h1>
<h2>Reunião geral (all-hands)</h2>
<p>A reunião geral do trimestre será no dia 30 de setembro, às 16h, no auditório e online.</p>
<h2>Semana da LGPD</h2>
<p>De 5 a 9 de outubro acontece a Semana da LGPD. O treinamento de 2 horas é obrigatório para todos os colaboradores.</p>
<h2>Plano Enterprise</h2>
<p>O plano Enterprise foi lançado em junho e já conta com 4 clientes ativos. Parabéns ao time comercial!</p>
<h2>Recesso de fim de ano</h2>
<p>A empresa entra em recesso de 24 de dezembro a 2 de janeiro. O plantão de suporte crítico continua ativo.</p>
</body></html>
""", encoding="utf-8")

print("Documentos gerados em", OUT)
for p in sorted(OUT.rglob("*")):
    if p.is_file():
        print(" -", p.relative_to(OUT))
